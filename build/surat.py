"""Surat Pernyataan Orisinalitas Karya (DOCX + PDF) sesuai ketentuan guidebook Hackvidia 11.0.

Catatan: guidebook menautkan template resmi (s.hmif.dev/TemplateOrisinalitasHackvidia11.0). Jika template resmi
bisa dibuka, salin isi pernyataan ini ke template tersebut. Ketua tanda tangan basah di atas meterai Rp10.000
(FAQ: meterai digital diperbolehkan); anggota lain boleh tanda tangan digital.
"""
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt

from common import ROOT, docx_to_pdf, load_config

FONT = "Times New Roman"


def run(par, text, size=12, bold=False):
    r = par.add_run(text)
    r.font.name = FONT; r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    r.font.size = Pt(size); r.bold = bold
    return r


def para(d, text="", size=12, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=6, spacing=1.15):
    p = d.add_paragraph()
    if text:
        run(p, text, size, bold)
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = spacing
    return p


def build(cfg):
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(1))
    para(d, "SURAT PERNYATAAN ORISINALITAS KARYA", 14, True, WD_ALIGN_PARAGRAPH.CENTER, 0, 1.15)
    para(d, "HACKVIDIA - ARKAVIDIA 11.0", 12, True, WD_ALIGN_PARAGRAPH.CENTER, 8, 1.15)
    para(d, "Kami yang bertanda tangan di bawah ini:")
    for label, val in [("Nama Tim", cfg["nama_tim"]), ("Judul Karya", cfg["judul_proposal"])]:
        p = para(d, "", after=2, align=WD_ALIGN_PARAGRAPH.LEFT)
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(3.7))
        run(p, f"{label}\t: "); run(p, val, bold=True)
    t = d.add_table(rows=1, cols=5); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(["No", "Nama", "NIM", "Program Studi", "Perguruan Tinggi"]):
        c = t.rows[0].cells[i]; c.text = ""; run(c.paragraphs[0], h, bold=True)
    for i, m in enumerate(cfg["anggota"], 1):
        cells = t.add_row().cells
        for j, v in enumerate([str(i), f"{m['nama']} ({m['peran']})", m["nim"], m["prodi"], m["kampus"]]):
            cells[j].text = ""; run(cells[j].paragraphs[0], v)
    for row in t.rows:
        for j, w in enumerate([1.0, 5.2, 2.8, 3.3, 3.6]):
            row.cells[j].width = Cm(w)
    para(d, "", after=2)
    para(d, "dengan ini menyatakan bahwa:")
    items = [
        "Karya yang kami ajukan pada kompetisi Hackvidia Arkavidia 11.0 merupakan karya orisinal tim kami dan dikembangkan khusus untuk kompetisi ini.",
        "Karya tersebut bukan hasil plagiarisme, belum pernah diikutsertakan dalam kompetisi lain, dan belum pernah memenangkan kompetisi sejenis.",
        "Seluruh data, kutipan, pustaka, aset, serta API/model pihak ketiga yang digunakan telah kami cantumkan sumbernya dan digunakan sesuai lisensinya.",
        "Kami tidak bekerja sama dengan tim lain dan tidak menerima bantuan pengerjaan karya dari pihak di luar tim, kecuali konsultasi dengan mentor yang disediakan panitia.",
        "Apabila di kemudian hari pernyataan ini terbukti tidak benar, kami bersedia menerima sanksi dari Panitia Arkavidia 11.0, termasuk diskualifikasi dan pencabutan gelar juara.",
    ]
    for i, it in enumerate(items, 1):
        p = para(d, "", after=3)
        p.paragraph_format.left_indent = Cm(0.8); p.paragraph_format.first_line_indent = Cm(-0.6)
        run(p, f"{i}.  {it}")
    para(d, "Demikian surat pernyataan ini kami buat dengan sebenar-benarnya tanpa paksaan dari pihak mana pun.", after=4)
    para(d, f"{cfg['kota_ttd']}, {cfg['tanggal_ttd']}", align=WD_ALIGN_PARAGRAPH.RIGHT, after=8)
    sig = d.add_table(rows=2, cols=2); sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    members = cfg["anggota"]
    blocks = []
    lead = members[0]
    blocks.append(("Ketua Tim,", "(meterai Rp10.000)\n\n\n", lead["nama"]))
    for m in members[1:]:
        blocks.append(("Anggota,", "\n\n\n", m["nama"]))
    # layout: 2 kolom x 2 baris
    cells = [sig.rows[0].cells[0], sig.rows[0].cells[1], sig.rows[1].cells[0], sig.rows[1].cells[1]]
    for c, (title, gap, name) in zip(cells, blocks):
        c.text = ""
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run(p, title + "\n" + gap); run(p, name, bold=True)
        p.paragraph_format.space_after = Pt(4); p.paragraph_format.line_spacing = 1.0
    return d


def main():
    cfg = load_config()
    d = build(cfg)
    out = ROOT / "05_Surat_Orisinalitas"
    docx_path = out / f"{cfg['nama_tim']}_SuratOrisinalitas.docx"
    d.save(docx_path)
    docx_to_pdf(docx_path, out / f"{cfg['nama_tim']}_SuratOrisinalitas.pdf")
    print("surat ok")


if __name__ == "__main__":
    main()
