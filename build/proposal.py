"""Proposal Hackvidia (DOCX + PDF), format: A4, Times New Roman 12, spasi 1,5, margin 1 inci.

Maks. 10 halaman teks (di luar cover dan lampiran wireframe; Daftar Pustaka ikut dihitung agar aman).
Konten mengikuti rubrik penilaian babak penyisihan Hackvidia Arkavidia 11.0.

Sitasi ditulis sebagai [@kunci] atau [@a;@b] dan dinomori otomatis menurut urutan kemunculan pertama.
Referensi yang tidak disitasi tidak dicetak.
"""
import math
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image

from common import BUILD, ROOT, docx_to_pdf, load_config

FONT = "Times New Roman"

# Semua rujukan dicek 23 Okt 2026 (lihat 02_Proposal/CATATAN_CEK_FAKTA.md).
REFS = {
    "iasc": "Kontan (6 Okt 2026). IASC terima 702.342 laporan kasus penipuan hingga September 2026. keuangan.kontan.co.id/news/iasc-terima-702342-laporan-kasus-penipuan-hingga-september-2026",
    "ojk": "OJK (2026). Siaran Pers SP 15/GKPB/OJK/I/2026: IASC berhasil kembalikan Rp161 miliar dana masyarakat korban scam. ojk.go.id",
    "jpnn": "JPNN (21 Agu 2026). Jawa Barat tertinggi kasus scam di Indonesia, 3 modus ini paling banyak. jpnn.com/news/jawa-barat-tertinggi-kasus-scam-di-indonesia-3-modus-ini-paling-banyak",
    "ai1379": "Kontan (2026). Per Juni 2026, IASC terima 1.379 laporan penipuan terindikasi memanfaatkan AI. keuangan.kontan.co.id",
    "komdigi": "Kompas (2 Jul 2026). Komdigi sebut banyak lansia jadi korban scam AI yang tiru suara pejabat. nasional.kompas.com/read/2026/07/02/20330831",
    "mcafee": "McAfee (1 Mei 2023). The Artificial Imposter (siaran pers). businesswire.com/news/home/20230501005587/en",
    "bps": "BPS (5 Mei 2026). Hasil SUPAS 2025: persentase lansia 11,97 persen. bps.go.id/id/pressrelease/2026/05/05/2645",
    "goodstats": "GoodStats (16 Jan 2026). Akses internet lansia capai 34% pada 2025 (mengolah BPS). data.goodstats.id",
    "snlik": "OJK & BPS (2025). Siaran Pers SP 69/OJK/GKPB/V/2025: Hasil SNLIK 2025. ojk.go.id",
    "james": "James, B. D., Boyle, P. A., & Bennett, D. A. (2014). Correlates of susceptibility to scams in older adults without dementia. J. Elder Abuse & Neglect, 26(2), 107-122.",
    "gchat": "Android Authority (11 Nov 2025). Pixel phones are now snitching on scammers in your chat apps. androidauthority.com/pixel-chat-app-scam-detection-3614116",
    "gfeb": "Google (25 Feb 2026). Staying one step ahead: strengthening Android's lead in scam protection. blog.google/security",
    "goct": "Google (14 Okt 2026). New Android protections to help keep you safe from mobile scams. blog.google/security/new-android-scam-protections",
    "truecaller": "Reuters via Malay Mail (8 Mei 2023). Truecaller aims to help WhatsApp users combat spam. malaymail.com",
    "aduan": "detikInet (16 Jul 2025). Cara laporkan penipuan online ke polisi dan Komdigi. inet.detik.com/security/d-8013809",
    "czaja": "Czaja, S. J., Boot, W. R., Charness, N., & Rogers, W. A. (2019). Designing for Older Adults (ed. 3). CRC Press.",
    "akhawe": "Akhawe, D., & Felt, A. P. (2013). Alice in Warningland: A large-scale field study of browser security warning effectiveness. USENIX Security 2013.",
    "anderson": "Anderson, B. B., dkk. (2015). How polymorphic warnings reduce habituation in the brain: Insights from an fMRI study. CHI 2015, 2883-2892.",
    "amershi": "Amershi, S., dkk. (2019). Guidelines for human-AI interaction. CHI 2019. doi:10.1145/3290605.3300233",
    "finra": "FINRA (2017). Regulatory Notice 17-11: Rule 2165 dan Rule 4512 (trusted contact person). finra.org/rules-guidance/notices/17-11",
    "roozenbeek": "Roozenbeek, J., dkk. (2022). Psychological inoculation improves resilience against misinformation on social media. Science Advances, 8(34), eabo6254.",
    "msspeech": "Microsoft. Azure AI Speech: language and voice support (su-ID, jv-ID). learn.microsoft.com",
    "otp": "Android Developers. Android 15 behavior changes: OTP redaction. developer.android.com/about/versions/15/behavior-changes-all",
    "csapi": "Android Developers. CallScreeningService. developer.android.com/reference/android/telecom/CallScreeningService",
    "wa": "WhatsApp Help Center. Tentang enkripsi end-to-end. faq.whatsapp.com/820124435853543",
    "restricted": "Google. Learn about restricted settings (Android Help). support.google.com/android/answer/12623953",
    "pdp": "UU No. 27 Tahun 2022 tentang Pelindungan Data Pribadi. peraturan.bpk.go.id/Details/229798",
    "sms": "Dwiyansaputra, R., dkk. (2021). Deteksi SMS spam berbahasa Indonesia menggunakan TF-IDF dan SGD classifier. JTIKA, 3(2) (data 1.143 SMS dari Rahmi & Wibisono, 2016).",
    "indonlu": "Wilie, B., dkk. (2020). IndoNLU: Benchmark and resources for evaluating Indonesian NLU. AACL-IJCNLP 2020, 843-857.",
    "statcounter": "StatCounter. Mobile OS market share Indonesia (Agu 2026). gs.statcounter.com/os-market-share/mobile/indonesia",
    "gemini": "Google AI for Developers. Gemini API pricing, diakses 23 Okt 2026. ai.google.dev/gemini-api/docs/pricing",
    "pojk": "OJK (2024). POJK No. 12 Tahun 2024 tentang Penerapan Strategi Anti Fraud bagi Lembaga Jasa Keuangan.",
    "taspen": "Bisnis.com (20 Mei 2026). PT Taspen minta pensiunan autentikasi (46 mitra bayar). finansial.bisnis.com/read/20260520/215/1975305",
    "psr": "Payment Systems Regulator UK (2024). PSR confirms its decision on APP scams reimbursement. psr.org.uk",
    "srf": "Baker McKenzie (2024). Singapore: Shared Responsibility Framework from 16 December 2024. insightplus.bakermckenzie.com",
}
_ORDER = []


def _cite(m):
    nums = []
    for key in m.group(1).split(";"):
        key = key.strip().lstrip("@")
        if key not in REFS:
            raise KeyError(f"rujukan tidak dikenal: {key}")
        if key not in _ORDER:
            _ORDER.append(key)
        nums.append(_ORDER.index(key) + 1)
    return "[" + "][".join(str(n) for n in nums) + "]"


def cite(text):
    return re.sub(r"\[(@[\w]+(?:;\s*@[\w]+)*)\]", _cite, text)


# ---------------------------------------------------------------- helpers
def set_cell_shading(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def set_run_font(run, size=12, bold=None, italic=None, color=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_rich(par, text, size=12, base_bold=False):
    """Mini-markup: **bold**, *italic*, sitasi [@kunci]."""
    text = cite(text)
    tokens = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text)
    for t in tokens:
        if not t:
            continue
        if t.startswith("**"):
            r = par.add_run(t[2:-2]); set_run_font(r, size, bold=True)
        elif t.startswith("*"):
            r = par.add_run(t[1:-1]); set_run_font(r, size, bold=base_bold, italic=True)
        else:
            r = par.add_run(t); set_run_font(r, size, bold=base_bold)
    return par


def fmt(par, spacing=1.5, after=6, before=0, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=None, keep_next=False):
    pf = par.paragraph_format
    pf.line_spacing = spacing
    pf.space_after = Pt(after)
    pf.space_before = Pt(before)
    par.alignment = align
    if indent is not None:
        pf.left_indent = indent
    pf.keep_with_next = keep_next
    return par


class Doc:
    def __init__(self):
        self.d = Document()
        st = self.d.styles["Normal"]
        st.font.name = FONT
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        st.font.size = Pt(12)
        sec = self.d.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, side, Inches(1))
        self.fig_no = 0
        self.tab_no = 0

    def p(self, text, **kw):
        par = self.d.add_paragraph()
        add_rich(par, text)
        return fmt(par, **kw)

    def h1(self, text):
        par = self.d.add_paragraph()
        r = par.add_run(text.upper()); set_run_font(r, 12, bold=True)
        return fmt(par, after=2, before=6, align=WD_ALIGN_PARAGRAPH.LEFT, keep_next=True)

    def h2(self, text):
        par = self.d.add_paragraph()
        r = par.add_run(text); set_run_font(r, 12, bold=True)
        return fmt(par, after=0, before=2, align=WD_ALIGN_PARAGRAPH.LEFT, keep_next=True)

    def bullets(self, items, numbered=False):
        for i, it in enumerate(items):
            par = self.d.add_paragraph()
            mark = f"{i + 1}." if numbered else "•"
            add_rich(par, f"{mark}\t{it}")
            pf = par.paragraph_format
            pf.left_indent = Cm(0.75)
            pf.first_line_indent = Cm(-0.6)
            pf.tab_stops.add_tab_stop(Cm(0.75))
            fmt(par, after=2)

    def caption(self, text):
        par = self.d.add_paragraph()
        add_rich(par, text, size=12)
        for r in par.runs:
            r.italic = True
        return fmt(par, spacing=1.0, after=6, align=WD_ALIGN_PARAGRAPH.CENTER)

    def table(self, header, rows, widths, title):
        self.tab_no += 1
        cap = self.d.add_paragraph()
        add_rich(cap, f"**Tabel {self.tab_no}.** {title}")
        fmt(cap, spacing=1.0, after=3, before=4, align=WD_ALIGN_PARAGRAPH.LEFT, keep_next=True)
        t = self.d.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.autofit = False
        tblPr = t._tbl.tblPr
        lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
        mar = OxmlElement("w:tblCellMar")
        for side in ("left", "right"):
            el = OxmlElement(f"w:{side}"); el.set(qn("w:w"), "70"); el.set(qn("w:type"), "dxa"); mar.append(el)
        tblPr.append(mar)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(header):
            c = t.rows[0].cells[i]
            c.text = ""
            par = c.paragraphs[0]
            r = par.add_run(h); set_run_font(r, 12, bold=True, color="FFFFFF")
            fmt(par, spacing=1.0, after=0, align=WD_ALIGN_PARAGRAPH.LEFT, keep_next=True)  # header tidak yatim
            set_cell_shading(c, "23395D")
        hdr_trPr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader"); hdr_trPr.append(th)
        for row in rows:
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = ""
                par = cells[i].paragraphs[0]
                add_rich(par, val)
                fmt(par, spacing=1.0, after=0, align=WD_ALIGN_PARAGRAPH.LEFT)
        for i, w in enumerate(widths):
            t.columns[i].width = Cm(w)
        for row in t.rows:
            trPr = row._tr.get_or_add_trPr()
            cs = OxmlElement("w:cantSplit"); trPr.append(cs)
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
        sp = self.d.add_paragraph(); fmt(sp, spacing=1.0, after=0)
        sp.paragraph_format.line_spacing = Pt(6)
        return t

    def figure(self, path, width_in, caption):
        self.fig_no += 1
        par = self.d.add_paragraph()
        par.add_run().add_picture(str(path), width=Inches(width_in))
        fmt(par, spacing=1.0, after=2, align=WD_ALIGN_PARAGRAPH.CENTER, keep_next=True)
        self.caption(f"Gambar {self.fig_no}. {caption}")

    def page_break(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_page_number_footer(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    par = footer.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = par.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag); run._r.append(el)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text; run._r.append(el)
    set_run_font(run, 12)


def restart_numbering(section, start=1):
    sectPr = section._sectPr
    pg = OxmlElement("w:pgNumType"); pg.set(qn("w:start"), str(start)); sectPr.append(pg)


def trim(path):
    im = Image.open(path).convert("RGB")
    from PIL import ImageChops
    bg = Image.new("RGB", im.size, (255, 255, 255))
    bbox = ImageChops.difference(im, bg).getbbox()
    out = BUILD / ("trim_" + Path(path).name)
    im.crop(bbox).save(out)
    return out


# ---------------------------------------------------------------- validasi (tidak pernah mengarang angka)
def wilson(pct, n, z=1.96):
    """Selang kepercayaan 95% Wilson untuk proporsi (pct dalam persen). Dibulatkan ke persen bulat."""
    k = round(pct * n / 100)
    p = k / n
    den = 1 + z * z / n
    c = (p + z * z / (2 * n)) / den
    h = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / den
    return round(100 * max(0, c - h)), round(100 * min(1, c + h))


def _pct(v, key, nkey):
    if v.get(key) is None or not v.get(nkey):
        return None
    lo, hi = wilson(v[key], v[nkey])
    return f"{v[key]}% (SK95% {lo}-{hi}%)"


def validation_paragraph(v):
    """Ringkasan hasil validasi untuk teks (dipakai juga oleh 06_Riset_Pengguna/olah_hasil.py).

    Setiap kalimat hanya muncul bila datanya ada di config_tim.json -> validasi.
    """
    parts = []
    if v.get("survei_n"):
        s = (f"Survei kepada {v['survei_n']} anak dewasa yang memiliki orang tua lansia: "
             f"{_pct(v, 'survei_pct_ortu_pernah_jadi_target', 'survei_n')} melaporkan orang tuanya pernah menjadi sasaran "
             f"penipuan dalam 12 bulan terakhir, {_pct(v, 'survei_pct_khawatir', 'survei_n')} khawatir orang tuanya tertipu, "
             f"dan {_pct(v, 'survei_pct_mau_jadi_pendamping', 'survei_n')} bersedia menjadi Pendamping")
        if v.get("survei_pct_bayar_min15rb") is not None:
            s += f"; {_pct(v, 'survei_pct_bayar_min15rb', 'survei_n')} menilai wajar membayar Rp15.000 atau lebih per bulan"
        parts.append(s + ".")
    if v.get("wawancara_lansia_n"):
        if v.get("wawancara_temuan"):
            parts.append(f"Wawancara dengan {v['wawancara_lansia_n']} lansia di Bandung: {v['wawancara_temuan']}.")
        else:
            parts.append(f"Kami mewawancarai {v['wawancara_lansia_n']} lansia di Bandung.")
    if v.get("uji_wireframe_n"):
        k = round(v['uji_wireframe_pct_paham'] * v['uji_wireframe_n'] / 100)
        parts.append(f"Uji wireframe: {k} dari {v['uji_wireframe_n']} lansia memahami peringatan dan memilih tindakan "
                     "aman tanpa bantuan.")
    if v.get("kutipan_lansia"):
        parts.append(f"Salah satu responden lansia menyampaikan, *\"{v['kutipan_lansia']}\"*")
    if parts:
        return "**Hasil validasi awal.** " + " ".join(parts)
    return ("**Status.** Protokol dan kriteria keputusan ditetapkan 23 Oktober 2026, sebelum data dikumpulkan; hasil "
            "dimasukkan apa adanya, termasuk bila target sampel tidak tercapai.")


def validation_rows(v):
    def hasil(keys_present, text_fn):
        return text_fn() if keys_present else "Dijadwalkan s.d. 25 Nov 2026"
    s_ok = bool(v.get("survei_n"))
    w_ok = bool(v.get("wawancara_lansia_n"))
    u_ok = bool(v.get("uji_wireframe_n"))
    def s1():
        pend = v.get("survei_pct_mau_jadi_pendamping")
        out = f"n={v['survei_n']}; Pendamping {pend}%"
        out += " → inti tetap" if pend >= 60 else (" → Pendamping opsional" if pend < 40 else " → zona tengah, diuji di pilot")
        bayar = v.get("survei_pct_bayar_min15rb")
        if bayar is not None:
            out += f"; bayar {bayar}%" + (" → premium dihapus" if bayar < 30 else " → premium diuji")
        return out
    def s3():
        pct = v["uji_wireframe_pct_paham"]
        return f"{round(pct * v['uji_wireframe_n'] / 100)}/{v['uji_wireframe_n']} paham" + (" → lanjut" if pct >= 80 else " → revisi A3")
    s_hasil = hasil(s_ok, s1)
    w_hasil = hasil(w_ok, lambda: f"n={v['wawancara_lansia_n']}; lihat teks")
    u_hasil = hasil(u_ok, s3)
    return [
        ["S1. Survei daring anak dewasa (target n≥30; orang tua/mertua 60+)",
         "% orang tua pernah jadi sasaran; % bersedia jadi Pendamping; % mau membayar ≥Rp15.000/bulan",
         "Pendamping ≥60%: Aturan Dua Orang tetap inti. <40%: Pendamping opsional, kader komunitas jadi bawaan. "
         "Bayar ≥Rp15 rb <30%: Paket Keluarga premium dihapus, fokus B2B2C.", s_hasil],
        ["S2. Wawancara semi-terstruktur lansia (8-10; posyandu lansia, pengajian, Bandung)",
         "Tema: panik, sungkan bertanya, preferensi suara/bahasa, batas privasi terhadap anak (analisis tematik, 2 pengode)",
         "Bila ≥50% menolak anak diberi tahu: ringkasan hanya dikirim setelah lansia menekan tombol (bawaan) dan opsi kader ditonjolkan.", w_hasil],
        ["S3. Uji wireframe A3/A5/A8 (n=5, *think-aloud*, tanpa bantuan)",
         "Paham = memilih tindakan aman (T1) dan menyebut ≥2 tanda bahaya (T2, *teach-back*)",
         "≥80% paham (≥4/5): lanjut. Kurang: revisi A3 dan uji ulang dengan 5 lansia baru sebelum final.", u_hasil],
    ]


# ---------------------------------------------------------------- content
def build(cfg):
    _ORDER.clear()
    D = Doc()
    d = D.d

    # ---------------- COVER (tanpa nomor halaman)
    for _ in range(2):
        fmt(d.add_paragraph(), spacing=1.0, after=0)
    for text, size, bold in [("PROPOSAL HACKVIDIA", 14, True), ("ARKAVIDIA 11.0", 14, True)]:
        par = d.add_paragraph(); r = par.add_run(text); set_run_font(r, size, bold=bold)
        fmt(par, spacing=1.15, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    fmt(d.add_paragraph(), spacing=1.0, after=18)
    par = d.add_paragraph(); r = par.add_run(cfg["nama_solusi"]); set_run_font(r, 30, bold=True, color="23395D")
    fmt(par, spacing=1.0, after=6, align=WD_ALIGN_PARAGRAPH.CENTER)
    par = d.add_paragraph(); r = par.add_run("Aturan Dua Orang Berbasis AI untuk Melindungi Lansia dari Penipuan Digital")
    set_run_font(r, 16, bold=True)
    fmt(par, spacing=1.15, after=18, align=WD_ALIGN_PARAGRAPH.CENTER)
    for line in ["Tema: Human-Centered Engineering in the AI Era", f"Subtema: {cfg['subtema']}"]:
        par = d.add_paragraph(); r = par.add_run(line); set_run_font(r, 12, italic=True)
        fmt(par, spacing=1.15, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    fmt(d.add_paragraph(), spacing=1.0, after=40)
    par = d.add_paragraph(); r = par.add_run(f"Disusun oleh Tim {cfg['nama_tim']}"); set_run_font(r, 13, bold=True)
    fmt(par, spacing=1.15, after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    for m in cfg["anggota"]:
        par = d.add_paragraph()
        r = par.add_run(f"{m['nama']} ({m['peran']}) - {m['prodi']}, {m['kampus']}"); set_run_font(r, 12)
        fmt(par, spacing=1.15, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    fmt(d.add_paragraph(), spacing=1.0, after=60)
    for line in ["Hackvidia - Arkavidia 11.0", "Himpunan Mahasiswa Informatika Institut Teknologi Bandung", "2026"]:
        par = d.add_paragraph(); r = par.add_run(line); set_run_font(r, 12, bold=True)
        fmt(par, spacing=1.15, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)

    sec = d.add_section(WD_SECTION.NEW_PAGE)
    restart_numbering(sec, 1)
    add_page_number_footer(sec)

    v = cfg["validasi"]

    # ---------------- 1
    D.h1("1. Ringkasan Eksekutif")
    D.p("Indonesia Anti-Scam Centre (IASC) menerima **702.342 laporan penipuan** dalam 22 bulan [@iasc]. Modus *fake call* "
        "(penipu menyamar sebagai keluarga, polisi, atau petugas) menimbulkan kerugian terlapor terbesar, **Rp2,09 triliun**, "
        "atau rata-rata **Rp32,9 juta per laporan**, dua kali lipat modus belanja daring [@jpnn]. Sasaran yang paling rentan "
        "terus bertambah: sekitar 11,6 juta lansia sudah daring [@bps;@goodstats], sementara literasi keuangan usia 51-79 tahun "
        "hanya 54,55% [@snlik]. Penipu menang ketika korbannya **panik dan sendirian**, dan suara keluarga kini dapat ditiru AI [@mcafee].")
    D.p("**Ide inti:** deteksi penipuan sedang menjadi fitur sistem operasi [@goct], tetapi belum ada yang menghadirkan "
        "**orang kedua yang dipercaya tepat pada menit kritis**. JagaOrtu membawa *aturan dua orang* (maker-checker) perbankan "
        "ke keluarga: saat pesan atau telepon menunjukkan tanda penipuan, JagaOrtu memberi jeda dan menjelaskan tanda bahayanya "
        "dengan suara dalam bahasa lansia, menghadirkan Pendamping keluarga dalam satu ketukan, dan memandu verifikasi dengan "
        "kata sandi keluarga yang tetap ampuh walau suara dikloning. AI dipakai secukupnya dalam empat lapis berambang jelas, "
        "dan setiap keputusan uang tetap di tangan manusia. Pembayar utama adalah bank/BPR penyalur pensiun yang wajib "
        "menjalankan strategi anti-fraud [@pojk]; biaya model bahasa ±Rp90 per pengguna per bulan (Tabel 9). Validasi disusun "
        "sebagai protokol pra-registrasi dengan kriteria keputusan yang ditetapkan sebelum data dikumpulkan (Tabel 2).")

    # ---------------- 2
    D.h1("2. Latar Belakang dan Analisis Masalah")
    D.h2("2.1 Skala dan urgensi")
    D.p("IASC mencatat kerugian Rp9,1 triliun dari 432.637 aduan hingga 14 Januari 2026 [@ojk] (rata-rata Rp21 juta). "
        "Pada November 2024-Juni 2026, tiga modus teratas adalah belanja daring (113.520 laporan; Rp1,78 T; rata-rata Rp15,7 "
        "juta), *fake call* (63.517; Rp2,09 T; Rp32,9 juta), dan investasi (32.251; Rp2,08 T; Rp64,5 juta); Jawa Barat "
        "tertinggi dengan 131.445 laporan [@jpnn]. Laporan terindikasi AI bertambah dari 13 (akhir 2024) menjadi total 1.379 "
        "(Juni 2026) [@ai1379], dan Komdigi menyebut banyak lansia menjadi korban peniruan suara pejabat [@komdigi]. Lansia "
        "mencapai 11,97% dari 284,67 juta penduduk [@bps] (±34,1 juta); 52,23% memakai ponsel dan 34,13% mengakses internet "
        "[@goodstats] (±11,6 juta). Rata-rata dan perkalian adalah perhitungan tim.")
    D.h2("2.2 Akar masalah: mengapa lansia tertipu")
    D.p("Modus yang dilaporkan berulang pada tiga mekanisme: **urgensi dan otoritas** (\"anak kecelakaan\", \"rekening akan "
        "diblokir\"), **isolasi** (\"jangan bilang siapa-siapa\"), dan **runtuhnya sinyal kepercayaan lama** karena suara dan "
        "nomor dapat dipalsukan [@mcafee]. Studi pada 639 lansia tanpa demensia menemukan kerentanan terhadap penipuan berkaitan "
        "dengan kognisi, kesejahteraan psikologis, dan literasi yang lebih rendah, bukan sekadar kurang hati-hati [@james]. "
        "Lansia butuh **jeda dan orang kedua yang dipercaya pada momen kritis**, bukan sekadar informasi tambahan. Persona "
        "kami: **Bu Rukmini (67)**, pensiunan guru di Bandung yang tinggal sendiri, mudah panik, dan sungkan bertanya; serta "
        "**Rina (34)**, anaknya di Jakarta, yang ingin melindungi tanpa memata-matai.")
    D.h2("2.3 Posisi terhadap solusi yang ada")
    D.table(["Solusi", "Yang sudah baik (kami akui)", "Batas bagi lansia Indonesia"], [
        ["Google Scam Detection (Messages, notifikasi chat, telepon)",
         "Di perangkat, gratis; sudah menilai notifikasi chat termasuk WhatsApp (Pixel 6+, AS) [@gchat]; akan dibawa ke "
         "12 merek Android, termasuk Xiaomi, OPPO, vivo [@goct]",
         "Sumber resmi belum menyebut Indonesia atau bahasa Indonesia/daerah [@gfeb;@goct]; hanya satu pengguna, tanpa orang kedua"],
        ["Truecaller, Getcontact", "Basis nomor spam komunitas; ID penelepon WhatsApp (beta 2023) [@truecaller]",
         "Bertumpu pada reputasi nomor (nomor baru lolos); tidak menilai isi permintaan"],
        ["IASC, aduannomor.id, cekrekening.id", "Jalur resmi lapor, blokir rekening/nomor [@ojk;@aduan]",
         "Setelah kejadian atau dibuka manual"],
        ["Peringatan aplikasi bank", "Muncul saat transaksi", "Teks umum; bank tidak melihat bujukan penipu di WhatsApp"],
    ], [3.6, 6.0, 6.3], "Perbandingan dengan solusi yang ada (dicek 23 Oktober 2026)")
    D.p("**Di mana kami kalah.** Pada deteksi, Google unggul: gratis, terpasang, dan meluas lewat produsen ponsel. JagaOrtu "
        "**tidak bersaing di deteksi**, tetapi membangun lapisan manusia di atasnya: Pendamping pada menit kritis, penjelasan "
        "bersuara berbahasa ibu, verifikasi prosedural, dan \"Jeda Transfer\" dengan bank. Sinyal deteksi OS kelak menjadi masukan L0.")
    D.h2("2.4 Validasi pengguna: protokol pra-registrasi")
    D.p(validation_paragraph(v))
    D.table(["Studi", "Ukuran utama", "Kriteria keputusan (sebelum data)", "Hasil"], validation_rows(v),
            [3.4, 4.4, 5.6, 2.5], "Protokol validasi (SK95% Wilson; persetujuan partisipan; tanpa nama, nomor, atau rekening)")

    # ---------------- 3
    D.h1("3. Deskripsi Solusi")
    D.p("JagaOrtu terdiri atas aplikasi Android untuk lansia dan aplikasi Pendamping untuk keluarga (Tabel 3, Gambar 1).")
    D.table(["Fitur", "Cara kerja", "Nilai bagi pengguna"], [
        ["**Penjaga Otomatis**", "Menilai notifikasi WhatsApp/SMS dari nomor tak dikenal atau berisi uang, tautan, file", "Tak perlu curiga lebih dulu"],
        ["**Tanya Dulu**", "Tombol besar untuk bercerita lewat suara atau membagikan pesan", "Bertanya terasa wajar"],
        ["**Aturan Dua Orang**", "Satu ketukan mengirim ringkasan risiko ke Pendamping; 10 menit tanpa respons → cadangan", "Tidak memutuskan sendirian"],
        ["**Verifikasi Keluarga**", "Telepon \"keluarga/polisi\": tanya kata sandi keluarga, tutup, telepon balik nomor tersimpan", "Ampuh walau suara dikloning"],
        ["**Latihan Waspada**", "Simulasi aman atas persetujuan lansia, maks. 1×/minggu, lalu pelajaran suara 30 detik", "Kewaspadaan terlatih"],
    ], [3.4, 8.6, 3.9], "Fitur inti JagaOrtu")
    D.figure(BUILD / "journey.png", 6.0, "Perjalanan pengguna tanpa dan dengan JagaOrtu")
    D.table(["Bukti", "Keputusan desain JagaOrtu"], [
        ["Kerentanan terkait kognisi dan literasi [@james]; lansia terbantu oleh beban memori dan pilihan yang kecil, target "
         "sentuh besar, umpan balik multimodal [@czaja]",
         "Penjelasan suara berbahasa ibu, maks. 3 tanda bahaya konkret, maks. 2 aksi per layar, teks ≥20 sp; dipasang oleh anak"],
        ["Kepatuhan pada peringatan sangat ditentukan desainnya [@akhawe]; peringatan yang bervariasi lebih tahan habituasi [@anderson]",
         "Dua tingkat (Kuning/Merah); \"anggaran alarm\" ≤1 Merah palsu per pengguna per minggu; isi spesifik per kasus"],
        ["AI perlu menjelaskan kemampuannya dan mudah dikoreksi [@amershi]",
         "Layar \"Yang dibaca JagaOrtu\" (A8); alasan selalu tampil; tombol \"Ini aman\" sebagai koreksi"],
        ["Regulasi pasar modal AS memakai *trusted contact person* untuk lansia [@finra]",
         "Aturan Dua Orang bersifat nasihat: Pendamping tak bisa menyetujui/memindahkan uang; lansia bebas mengganti Pendamping"],
    ], [7.4, 8.5], "Dari bukti riset ke keputusan desain")
    D.p("Prinsip yang dijaga: **martabat** (menjelaskan tanpa menghakimi), **otonomi** (JagaOrtu tidak pernah memblokir atau "
        "mentransfer), **bahasa ibu**, **kesederhanaan**, dan **privasi**. Latihan Waspada memakai prinsip inokulasi yang "
        "meningkatkan ketahanan terhadap manipulasi informasi [@roozenbeek]; efeknya pada penipuan finansial diukur di pilot (§8.2).")

    # ---------------- 4
    D.h1("4. Implementasi AI")
    D.h2("4.1 Empat lapis keputusan dan ambangnya")
    D.p("Tema Hackvidia menuntut kejelasan kapan AI diperlukan. Setiap lapis hanya dipakai bila lapis sebelumnya tidak yakin (Tabel 5).")
    D.table(["Lapis", "Aturan keputusan (ambang awal)", "Keluaran"], [
        ["**L0 Aturan pasti** (HP, <50 ms, tanpa AI)", "Aturan pasti (file .apk, minta OTP/PIN, nomor/rekening/domain terlapor) → **Merah**; "
         "sinyal lemah (tautan, \"nomor baru\", \"jangan bilang\") menjadi fitur L1", "Tanda bahaya"],
        ["**L1 Model di perangkat** (offline, <100 ms)", "Skor *s*<0,40 **Hijau** (diam); 0,40-0,70 → L2; *s*≥0,70 **Merah**", "Skor + tanda bahaya"],
        ["**L2 Model bahasa besar** (cloud, teks tersamar, <4 dtk)", "JSON tervalidasi skema: risiko, modus (8 kelas), ≤3 tanda bahaya, "
         "penjelasan ≤2 kalimat; tinggi → Merah, sedang → **Kuning**, rendah → Hijau; gagal/timeout → Kuning", "Penjelasan + saran"],
        ["**L3 Manusia**", "Merah: layar jeda + suara + \"Tanya [Pendamping] dulu\". Kuning: notifikasi bersuara; Pendamping "
         "diberi tahu hanya bila lansia meminta", "Keputusan uang"],
    ], [4.2, 9.0, 2.7], "Lapis keputusan, ambang, dan keluaran")
    D.p("Ambang awal dikalibrasi pada data validasi: ambang bawah agar recall kelas berisiko tinggi pada L1+L2 ≥0,95, ambang "
        "atas agar presisi Merah ≥0,80. Dua aturan **monoton** menahan manipulasi: L2 tidak pernah menurunkan putusan L0, dan "
        "untuk pesan berisi uang, tautan, atau file tidak dapat menurunkan risiko di bawah Kuning. **Domain AI:** NLP untuk teks "
        "campur bahasa; *speech* (STT untuk Tanya Dulu; TTS Indonesia di perangkat dan audio templat Sunda/Jawa yang dibuat "
        "sekali dengan suara neural [@msspeech]); OCR untuk tangkapan layar. JagaOrtu sengaja **tidak mengklaim mendeteksi "
        "deepfake suara**; verifikasi prosedural tetap efektif apa pun kualitas suaranya.")
    D.h2("4.2 Arsitektur: apa yang di perangkat, apa yang di cloud")
    D.figure(BUILD / "arch.png", 6.0, "Arsitektur JagaOrtu dan alur keputusan empat lapis")
    D.p("**Di perangkat:** pembacaan notifikasi atas izin eksplisit, L0, L1, OCR, penyamaran nomor/nama/rekening, dan suara. "
        "**Di cloud:** hanya L2 atas teks tersamar (tanpa penyimpanan isi), ringkasan ke Pendamping, dan pembaruan aturan "
        "mingguan. Tanpa internet, L0 dan L1 tetap melindungi.")
    D.h2("4.3 Evaluasi dan penanganan alarm palsu")
    D.table(["Ukuran", "Target", "Cara ukur"], [
        ["Recall kelas berisiko tinggi", "≥0,90", "Data uji ditahan, dipisah menurut waktu dan sumber; variasi sintetis tidak bocor ke uji"],
        ["Alarm palsu (pesan normal + promo)", "Kuning+Merah ≤10%; Merah ≤3%", "Promo sebagai negatif sulit; di pilot ≤1 Merah palsu/pengguna/minggu"],
        ["Porsi ke L2; latensi L0+L1", "≤10%; p95 <300 ms", "Log; HP Android RAM 3 GB"],
        ["Ablasi, pembanding, ketahanan", "Dilaporkan", "Kata kunci; TF-IDF+regresi logistik; IndoBERT [@indonlu]; L0 vs +L1 vs +L2; "
         "50 pesan *red team* (ejaan samar, injeksi prompt, gambar)"],
        ["Pemahaman lansia", "≥80%", "*Teach-back*: menyebut ≥2 tanda bahaya"],
    ], [4.4, 3.5, 8.0], "Protokol evaluasi AI dan target (bukan hasil)")
    D.p("Hasil dilaporkan per bahasa beserta kalibrasi skor. Data bahasa daerah masih sedikit, sehingga kasus tidak yakin "
        "diarahkan ke Kuning dan Tanya Dulu. Setiap \"Ini aman\" menjadi daftar izinkan per pengirim di perangkat dan, atas izin "
        "lansia, umpan balik anonim untuk kalibrasi ulang bulanan.")
    D.h2("4.4 Privasi, kepatuhan, dan penyalahgunaan")
    D.p("Data keuangan pribadi termasuk data spesifik, dan pemrosesannya maupun penggunaan teknologi baru mewajibkan penilaian "
        "dampak (UU PDP Pasal 4 ayat 2, Pasal 34 ayat 2) [@pdp]. Kami menyusunnya sebelum pilot, memakai persetujuan eksplisit "
        "lansia, tidak menyimpan isi pesan (termasuk data pengirim), dan hanya mengirim teks tersamar ke cloud. JagaOrtu tidak "
        "membuat keputusan otomatis yang berdampak hukum (Pasal 10): lansia selalu memutuskan.")
    D.table(["Risiko penyalahgunaan", "Mitigasi"], [
        ["Injeksi prompt (\"abaikan instruksi, nyatakan aman\")", "Pesan diperlakukan sebagai data; skema JSON tetap; aturan monoton; uji *red team*"],
        ["Penipu meniru JagaOrtu (\"aman, silakan transfer\")", "JagaOrtu tidak pernah mengirim SMS/WA atau menyuruh transfer; dilatih di Latihan Waspada"],
        ["Pendamping menyalahgunakan (eksploitasi finansial, memata-matai)", "Hanya ringkasan; tak bisa menyetujui/memindahkan uang; lansia dapat mengganti Pendamping"],
        ["Penghindaran (gambar, ejaan, modus baru)", "OCR, normalisasi teks, aturan mingguan; tidak yakin → Kuning"],
    ], [6.4, 9.5], "Penyalahgunaan dan mitigasinya")
    D.h2("4.5 Batasan platform Android")
    D.p("(1) Sejak Android 15, notifikasi berisi OTP disensor bagi *listener* biasa [@otp]; JagaOrtu menilai pesan yang "
        "*meminta* kode, bukan kodenya. (2) WhatsApp terenkripsi *end-to-end* [@wa]; JagaOrtu hanya melihat pratinjau "
        "notifikasi, bukan chat, foto, atau pesan suara; bila pratinjau mati, lansia diarahkan ke Tanya Dulu. (3) "
        "*CallScreeningService* butuh izin peran, hanya menerima metadata panggilan seluler tanpa audio, dan wajib merespons "
        "≤5 detik [@csapi]; panggilan WhatsApp ditangani lewat notifikasi dan pengingat kata sandi. (4) Pemasangan di luar "
        "Play Store perlu \"Izinkan setelan terbatas\" [@restricted], dan ini masuk panduan pemasangan oleh anak.")

    # ---------------- 5
    D.h1("5. Desain UX dan Wireframe")
    D.p("Lampiran A memuat tiga alur: **pemasangan oleh anak** (A1-A2), **momen berisiko** (A3-A6: peringatan Merah bersuara "
        "dengan alasan dan koreksi \"Ini aman\", Tanya Dulu, pengingat kata sandi saat telepon, ringkasan untuk Pendamping), dan "
        "**penguatan jangka panjang** (A7-A8: Latihan Waspada, kendali privasi dan Pendamping oleh lansia). Hasil S3 (Tabel 2) "
        "menentukan revisi A3 sebelum final.")

    # ---------------- 6
    D.h1("6. Rencana Data")
    D.p("**Sumber:** (1) 1.143 SMS berbahasa Indonesia (569 normal, 574 spam) [@sms], dilabel ulang ke taksonomi kami; (2) "
        "contoh modus dari peringatan resmi OJK, Komdigi, BSSN, bank, dan berita; (3) variasi sintetis termasuk Sunda/Jawa yang "
        "ditinjau penutur asli; (4) laporan pengguna opt-in dan anonim (dihapus setelah 30 hari kecuali disetujui). "
        "**Taksonomi:** delapan kelas selaras kategori IASC (keluarga darurat, menyamar petugas, tautan/file berbahaya, hadiah, "
        "investasi, salah transfer, belanja fiktif, pinjaman ilegal) ditambah aman dan promo. **Pelabelan:** dua anotator, "
        "adjudikasi, target Cohen's kappa ≥0,8.")

    # ---------------- 7
    D.h1("7. Rancangan Implementasi, Kelayakan, dan Skalabilitas")
    D.p("**Tech stack:** Kotlin, ML Kit OCR, TextToSpeech perangkat; L1 model linear TF-IDF di MVP, lalu TensorFlow Lite "
        "hasil distilasi di pilot; FastAPI serverless; PostgreSQL; API LLM berkeluaran JSON; Pendamping berupa aplikasi web "
        "progresif dengan *push*. Android dipakai ±77% pengguna ponsel di Indonesia [@statcounter].")
    D.table(["Fase", "Waktu", "Target"], [
        ["0. MVP", "Hacking Day", "Penjaga Otomatis WhatsApp/SMS, L0-L3 dengan log, peringatan bersuara, Pendamping + eskalasi"],
        ["1. Pilot", "Feb-Apr 2027", "30 keluarga di Bandung (§8.2) dan penilaian dampak pelindungan data"],
        ["2. Skala", "Mei 2027-2028", "1.000 keluarga lewat komunitas; API \"Jeda Transfer\" dengan bank/BPR penyalur pensiun"],
    ], [2.2, 3.2, 10.5], "Peta jalan pengembangan")
    D.p("**Kelayakan:** semua komponen memakai API Android resmi dan model yang tersedia, tanpa melatih model besar; karena "
        "sebagian besar keputusan diambil di perangkat, biaya cloud tetap rendah saat pengguna bertambah. **Risiko lain:** "
        "pemasangan (dipandu anak, tiga langkah), iOS belum didukung, dan lansia tanpa keluarga dekat (Pendamping dari kader).")

    # ---------------- 8
    D.h1("8. Model Bisnis, Dampak, dan Keberlanjutan")
    D.h2("8.1 Siapa membayar dan mengapa")
    D.p("**B2B2C (utama):** bank dan BPR wajib menerapkan strategi anti-fraud empat pilar yang cakupannya meliputi konsumen, "
        "termasuk sosialisasi kepada pihak eksternal [@pojk]. JagaOrtu menjadi komponen *pencegahan* bagi nasabah lansia dan "
        "memberi laporan agregat anonim untuk pilar *pemantauan*. Titik masuk: **bank penyalur pensiun**; Taspen bekerja sama "
        "dengan 46 mitra bayar [@taspen], sehingga nasabah lansia terkumpul di sedikit lembaga. Jujur saja, di Indonesia kerugian "
        "*scam* umumnya ditanggung korban, sehingga nilai bagi bank saat ini adalah kepatuhan, beban pengaduan, dan loyalitas "
        "pensiunan. Inggris mewajibkan bank mengganti korban penipuan transfer hingga £85.000 sejak Oktober 2024 [@psr] dan "
        "Singapura membagi tanggung jawab bank-operator untuk *phishing* [@srf]; bila arah ini diikuti, nilai \"Jeda Transfer\" "
        "bagi bank menjadi langsung. **Paket Keluarga:** dasar gratis; premium Rp25.000/bulan hanya bila lolos kriteria S1. "
        "**Kanal:** pilot komunitas gratis di Bandung → satu bank/BPR penyalur pensiun di Jawa Barat (*white-label*) → e-wallet "
        "(diatur Bank Indonesia) dan operator seluler.")
    D.table(["Per pengguna lansia per bulan", "Asumsi (diukur saat pilot)", "Nilai"], [
        ["Pendapatan B2B2C", "Lisensi (asumsi harga, belum diuji ke bank)", "Rp4.000"],
        ["LLM (L2)", "300 notifikasi dinilai; 10% ke L2 = 30 × (800 token masuk + 250 keluar); US$0,10/0,40 per 1 juta token [@gemini]; Rp16.500/US$", "±Rp90"],
        ["LLM skenario buruk", "Porsi L2 30%, token 2×", "±Rp540"],
        ["Suara; server dan *push*; dukungan dan kader", "TTS perangkat + audio templat; asumsi skala 10.000 pengguna", "Rp0; Rp300; Rp700"],
        ["**Margin kotor**", "Kasus dasar / skenario buruk", "**±73% / ±61%**"],
    ], [4.4, 9.0, 2.5], "Ekonomi unit (semua angka asumsi kecuali harga API)")
    D.h2("8.2 Dampak dan cara mengukurnya")
    D.p("Kami **tidak** memproyeksikan kerugian tercegah sebelum ada data pilot. Ukuran bertingkat: **keluaran** (momen "
        "berisiko tertangkap, Pendamping merespons ≤10 menit, alarm palsu), **hasil** (tingkat tindakan aman pada simulasi "
        "Latihan Waspada, rasa percaya diri lansia), dan **dampak** (kerugian yang dilaporkan keluarga; di fase mitra, data "
        "fraud agregat bank). Pilot 30 keluarga memakai desain daftar tunggu: 15 keluarga mulai di minggu 0 dan 15 di minggu 6, "
        "dengan ukuran utama selisih tingkat tindakan aman pada simulasi di minggu 6. Kerugian nyata terlalu jarang untuk diuji "
        "pada 30 keluarga; ini keterbatasan yang kami nyatakan. Sebagai pembanding, satu *fake call* yang tercegah menghindari "
        "rata-rata Rp32,9 juta [@jpnn], sehingga lisensi Rp48.000 per pengguna per tahun impas secara sosial bila mencegah satu "
        "kejadian per ±685 pengguna. Keberlanjutan dijaga oleh pendapatan B2B2C, efek jaringan keluarga, dan korpus modus lokal "
        "yang tumbuh. Tujuannya sederhana: tidak ada lansia yang menghadapi penipu sendirian.")

    # ---------------- Daftar Pustaka
    D.h1("Daftar Pustaka")
    for i, key in enumerate(_ORDER, 1):
        par = d.add_paragraph()
        r = par.add_run(f"[{i}] {REFS[key]}"); set_run_font(r, 12)
        pf = par.paragraph_format
        pf.left_indent = Cm(0.9); pf.first_line_indent = Cm(-0.9)
        fmt(par, spacing=1.0, after=2, align=WD_ALIGN_PARAGRAPH.LEFT)
    unused = [k for k in REFS if k not in _ORDER]
    if unused:
        print("  (rujukan tidak disitasi, tidak dicetak:", ", ".join(unused) + ")")

    # ---------------- Lampiran (tidak dihitung dalam batas 10 halaman)
    sec2 = d.add_section(WD_SECTION.NEW_PAGE)
    for el in sec2._sectPr.findall(qn("w:pgNumType")):
        sec2._sectPr.remove(el)
    for i, name in enumerate(["Lampiran_A1", "Lampiran_A2", "Lampiran_A3", "Lampiran_A4"]):
        if i == 0:
            D.h1("Lampiran A. Wireframe Low-Fidelity")
        img = trim(ROOT / "03_Wireframe" / f"{name}.png")
        w, h = Image.open(img).size
        width = min(6.27, 8.9 * w / h)
        par = d.add_paragraph(); par.add_run().add_picture(str(img), width=Inches(width))
        fmt(par, spacing=1.0, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
        if i < 3:
            D.page_break()
    return d


def check_pages(pdf_path, limit=10):
    """Hitung halaman teks utama (setelah cover, sebelum 'LAMPIRAN A') dan beri peringatan bila > limit."""
    try:
        from pypdf import PdfReader
    except ImportError:
        print("  (lewati cek halaman: pypdf belum terpasang)")
        return None
    pages = PdfReader(str(pdf_path)).pages
    first_lamp = next((i for i, pg in enumerate(pages) if "LAMPIRAN A" in (pg.extract_text() or "")), len(pages))
    n = first_lamp - 1  # halaman 0 = cover
    status = "OK" if n <= limit else "MELEBIHI BATAS"
    print(f"  halaman teks: {n} (batas {limit}) -> {status}; total PDF {len(pages)} halaman")
    return n


def main():
    cfg = load_config()
    d = build(cfg)
    out_dir = ROOT / "02_Proposal"
    docx_path = out_dir / f"{cfg['nama_tim']}_Proposal.docx"
    pdf_path = out_dir / f"{cfg['nama_tim']}_Proposal.pdf"
    d.save(docx_path)
    if docx_to_pdf(docx_path, pdf_path):
        print("proposal ok ->", pdf_path)
        n = check_pages(pdf_path)
        if n is not None and n > 10:
            raise SystemExit("Proposal melebihi 10 halaman teks. Ringkas isi sebelum dikirim.")


if __name__ == "__main__":
    main()
