"""Proposal Hackvidia (DOCX + PDF), format: A4, Times New Roman 12, spasi 1,5, margin 1 inci.

Maks. 10 halaman teks (di luar cover dan lampiran wireframe). Konten mengikuti rubrik penilaian
babak penyisihan Hackvidia Arkavidia 11.0.
"""
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image

from common import BUILD, ROOT, docx_to_pdf, load_config

FONT = "Times New Roman"
REFS = [
    "Kontan. (2026, 6 Oktober). IASC terima 702.342 laporan kasus penipuan hingga September 2026. https://keuangan.kontan.co.id/news/iasc-terima-702342-laporan-kasus-penipuan-hingga-september-2026",
    "Otoritas Jasa Keuangan. (2026). Siaran Pers SP 15/GKPB/OJK/I/2026 tentang capaian IASC. https://ojk.go.id/id/berita-dan-kegiatan/siaran-pers/Pages/IASC-Berhasil-Kembalikan-Rp161-Miliar-Dana-Masyarakat-Korban-Scam.aspx",
    "JPNN. (2026, 21 Agustus). Jawa Barat tertinggi kasus scam di Indonesia, 3 modus ini paling banyak. https://www.jpnn.com/news/jawa-barat-tertinggi-kasus-scam-di-indonesia-3-modus-ini-paling-banyak",
    "Kontan. (2026). Per Juni 2026, IASC terima 1.379 laporan penipuan terindikasi memanfaatkan AI. https://keuangan.kontan.co.id/news/per-juni-2026-iasc-terima-1379-laporan-penipuan-terindikasi-memanfaatkan-ai",
    "Kompas. (2026, 2 Juli). Komdigi sebut banyak lansia jadi korban scam AI yang tiru suara pejabat. https://nasional.kompas.com/read/2026/07/02/20330831/komdigi-sebut-banyak-lansia-jadi-korban-scam-ai-yang-tiru-suara-pejabat",
    "McAfee. (2023, 1 Mei). The Artificial Imposter: riset kloning suara AI dan survei 7.054 responden di tujuh negara (siaran pers). Business Wire. https://www.businesswire.com/news/home/20230501005587/en",
    "Badan Pusat Statistik. (2026, 5 Mei). SUPAS 2025: angka kelahiran total (TFR) 2,13, angka kematian bayi (IMR) 14,12, dan persentase lansia mencapai 11,97 persen (Berita Resmi Statistik; jumlah penduduk 284,67 juta). https://www.bps.go.id/id/pressrelease/2026/05/05/2645/hasil-survey-penduduk-antar-sensus.html",
    "GoodStats. (2026, 16 Januari). Akses internet lansia terus meningkat, capai 34% pada 2025 (mengolah BPS, Statistik Penduduk Lanjut Usia 2025). https://data.goodstats.id/statistic/akses-internet-lansia-terus-meningkat-capai-34-pada-2025-FxJIB",
    "Otoritas Jasa Keuangan & Badan Pusat Statistik. (2025). Siaran Pers SP 69: Hasil SNLIK 2025. https://ojk.go.id/id/berita-dan-kegiatan/siaran-pers/Pages/OJK-dan-BPS-Umumkan-Hasil-Survei-Nasional-Literasi-Dan-Inklusi-Keuangan-SNLIK-Tahun-2025.aspx",
    "Kompas Tekno. (2026, 4 Juni). Telepon scam makin canggih, Google rilis fitur penangkalnya. https://tekno.kompas.com/read/2026/06/04/08060077/telepon-scam-makin-canggih-google-rilis-fitur-penangkalnya",
    "Otoritas Jasa Keuangan. (2024). POJK Nomor 12 Tahun 2024 tentang Penerapan Strategi Anti Fraud bagi Lembaga Jasa Keuangan. https://ojk.go.id/en/regulasi/Pages/12-Tahun-2024-Penerapan-Strategi-Anti-Fraud-bagi-Lembaga-Jasa-Keuangan.aspx",
    "Republik Indonesia. (2022). Undang-Undang Nomor 27 Tahun 2022 tentang Pelindungan Data Pribadi. https://peraturan.bpk.go.id/Details/229798/uu-no-27-tahun-2022",
    "Federal Bureau of Investigation. (2025). Internet Crime Report 2024 (IC3). https://www.ic3.gov/AnnualReport/Reports/2024_IC3Report.pdf",
    "Android Developers. NotificationListenerService (API reference). https://developer.android.com/reference/android/service/notification/NotificationListenerService",
    "Rahmi, F., & Wibisono, Y. Dataset SMS spam berbahasa Indonesia (1.143 pesan), dirujuk dalam Jurnal JTIKA. https://jtika.if.unram.ac.id/index.php/JTIKA/article/download/145/83",
    "Microsoft. Azure AI Speech: language and voice support (id-ID, su-ID, jv-ID). https://learn.microsoft.com/azure/ai-services/speech-service/language-support",
    "Roozenbeek, J., van der Linden, S., Goldberg, B., Rathje, S., & Lewandowsky, S. (2022). Psychological inoculation improves resilience against misinformation on social media. Science Advances, 8(34), eabo6254.",
    "Android Developers. Behavior changes: all apps (Android 15), bagian Security and privacy: OTP redaction. https://developer.android.com/about/versions/15/behavior-changes-all",
    "Android Developers. CallScreeningService (API reference). https://developer.android.com/reference/android/telecom/CallScreeningService",
    "WhatsApp Help Center. Tentang enkripsi end-to-end. https://faq.whatsapp.com/820124435853543",
    "Google. Learn about restricted settings (Android Help). https://support.google.com/android/answer/12623953",
]


# ---------------------------------------------------------------- helpers
def set_cell_shading(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def set_run_font(run, size=12, bold=None, italic=None, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_rich(par, text, size=12, base_bold=False):
    """Mini-markup: **bold**, *italic*."""
    import re
    tokens = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text)
    for t in tokens:
        if not t:
            continue
        if t.startswith("**"):
            r = par.add_run(t[2:-2]); set_run_font(r, size, bold=True)
        elif t.startswith("*"):
            r = par.add_run(t[1:-1]); set_run_font(r, size, bold=base_bold, italic=True)
        else:
            r = par.add_run(t); set_run_font(r, size, bold=base_bold)
    return par


def fmt(par, spacing=1.5, after=6, before=0, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=None, keep_next=False):
    pf = par.paragraph_format
    pf.line_spacing = spacing
    pf.space_after = Pt(after)
    pf.space_before = Pt(before)
    par.alignment = align
    if indent is not None:
        pf.left_indent = indent
    pf.keep_with_next = keep_next
    return par


class Doc:
    def __init__(self):
        self.d = Document()
        st = self.d.styles["Normal"]
        st.font.name = FONT
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        st.font.size = Pt(12)
        sec = self.d.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, side, Inches(1))
        self.fig_no = 0
        self.tab_no = 0

    def p(self, text, **kw):
        par = self.d.add_paragraph()
        add_rich(par, text)
        return fmt(par, **kw)

    def h1(self, text):
        par = self.d.add_paragraph()
        r = par.add_run(text.upper()); set_run_font(r, 12, bold=True)
        return fmt(par, after=4, before=8, align=WD_ALIGN_PARAGRAPH.LEFT, keep_next=True)

    def h2(self, text):
        par = self.d.add_paragraph()
        r = par.add_run(text); set_run_font(r, 12, bold=True)
        return fmt(par, after=2, before=2, align=WD_ALIGN_PARAGRAPH.LEFT, keep_next=True)

    def bullets(self, items, numbered=False):
        for i, it in enumerate(items):
            par = self.d.add_paragraph()
            mark = f"{i + 1}." if numbered else "•"
            add_rich(par, f"{mark}\t{it}")
            pf = par.paragraph_format
            pf.left_indent = Cm(0.75)
            pf.first_line_indent = Cm(-0.6)
            pf.tab_stops.add_tab_stop(Cm(0.75))
            fmt(par, after=2)

    def caption(self, text):
        par = self.d.add_paragraph()
        add_rich(par, text, size=12)
        for r in par.runs:
            r.italic = True
        return fmt(par, spacing=1.0, after=8, align=WD_ALIGN_PARAGRAPH.CENTER)

    def table(self, header, rows, widths, title):
        self.tab_no += 1
        cap = self.d.add_paragraph()
        add_rich(cap, f"**Tabel {self.tab_no}.** {title}")
        fmt(cap, spacing=1.0, after=3, before=4, align=WD_ALIGN_PARAGRAPH.LEFT, keep_next=True)
        t = self.d.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.autofit = False
        tblPr = t._tbl.tblPr
        lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(header):
            c = t.rows[0].cells[i]
            c.text = ""
            par = c.paragraphs[0]
            r = par.add_run(h); set_run_font(r, 12, bold=True, color="FFFFFF")
            fmt(par, spacing=1.0, after=0, align=WD_ALIGN_PARAGRAPH.LEFT)
            set_cell_shading(c, "23395D")
        for row in rows:
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = ""
                par = cells[i].paragraphs[0]
                add_rich(par, val)
                fmt(par, spacing=1.0, after=0, align=WD_ALIGN_PARAGRAPH.LEFT)
        for i, w in enumerate(widths):
            t.columns[i].width = Cm(w)
        for row in t.rows:
            trPr = row._tr.get_or_add_trPr()
            cs = OxmlElement("w:cantSplit"); trPr.append(cs)
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
        sp = self.d.add_paragraph(); fmt(sp, spacing=1.0, after=2)
        return t

    def figure(self, path, width_in, caption):
        self.fig_no += 1
        par = self.d.add_paragraph()
        par.add_run().add_picture(str(path), width=Inches(width_in))
        fmt(par, spacing=1.0, after=2, align=WD_ALIGN_PARAGRAPH.CENTER, keep_next=True)
        self.caption(f"Gambar {self.fig_no}. {caption}")

    def page_break(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_page_number_footer(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    par = footer.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = par.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag); run._r.append(el)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text; run._r.append(el)
    set_run_font(run, 12)


def restart_numbering(section, start=1):
    sectPr = section._sectPr
    pg = OxmlElement("w:pgNumType"); pg.set(qn("w:start"), str(start)); sectPr.append(pg)


def trim(path):
    im = Image.open(path).convert("RGB")
    from PIL import ImageChops
    bg = Image.new("RGB", im.size, (255, 255, 255))
    bbox = ImageChops.difference(im, bg).getbbox()
    out = BUILD / ("trim_" + Path(path).name)
    im.crop(bbox).save(out)
    return out


# ---------------------------------------------------------------- content
def validation_paragraph(v):
    """Paragraf validasi dari config_tim.json -> validasi. Setiap bagian hanya muncul bila datanya ada (tidak mengarang).

    Angka diisi otomatis oleh 06_Riset_Pengguna/olah_hasil.py --tulis.
    """
    parts = []
    if v.get("survei_n"):
        parts.append(f"Survei kepada {v['survei_n']} anak dewasa yang memiliki orang tua lansia menunjukkan "
                     f"{v['survei_pct_ortu_pernah_jadi_target']}% orang tua mereka pernah menjadi sasaran pesan atau telepon "
                     f"penipuan dalam 12 bulan terakhir, {v['survei_pct_khawatir']}% khawatir orang tuanya tertipu, dan "
                     f"{v['survei_pct_mau_jadi_pendamping']}% bersedia menjadi Pendamping.")
    if v.get("wawancara_lansia_n"):
        if v.get("wawancara_temuan"):
            parts.append(f"Wawancara dengan {v['wawancara_lansia_n']} lansia di Bandung: {v['wawancara_temuan']}.")
        else:
            parts.append(f"Kami juga mewawancarai {v['wawancara_lansia_n']} lansia di Bandung.")
    if v.get("uji_wireframe_n"):
        parts.append(f"Uji wireframe dengan {v['uji_wireframe_n']} lansia: {v['uji_wireframe_pct_paham']}% memahami "
                     "peringatan dan memilih tindakan aman tanpa bantuan.")
    if v.get("kutipan_lansia"):
        parts.append(f"Salah satu responden lansia menyampaikan, *\"{v['kutipan_lansia']}\"*")
    if parts:
        return "**Validasi awal.** " + " ".join(parts)
    return ("**Validasi pengguna.** Desain ini divalidasi dalam tiga langkah sebelum Hacking Day: (1) survei daring kepada "
            "minimal 30 anak dewasa yang memiliki orang tua lansia, (2) wawancara mendalam dengan 8-10 lansia di komunitas "
            "posyandu lansia dan pengajian di Bandung, serta (3) uji keterbacaan wireframe dengan 5 lansia (tugas: memahami "
            "peringatan dan memilih tindakan aman). Hasilnya menjadi dasar penyesuaian desain dan ambang peringatan sebelum final.")


def build(cfg):
    D = Doc()
    d = D.d

    # ---------------- COVER (tanpa nomor halaman)
    for _ in range(2):
        fmt(d.add_paragraph(), spacing=1.0, after=0)
    for text, size, bold in [("PROPOSAL HACKVIDIA", 14, True), ("ARKAVIDIA 11.0", 14, True)]:
        par = d.add_paragraph(); r = par.add_run(text); set_run_font(r, size, bold=bold)
        fmt(par, spacing=1.15, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    fmt(d.add_paragraph(), spacing=1.0, after=18)
    par = d.add_paragraph(); r = par.add_run(cfg["nama_solusi"]); set_run_font(r, 30, bold=True, color="23395D")
    fmt(par, spacing=1.0, after=6, align=WD_ALIGN_PARAGRAPH.CENTER)
    par = d.add_paragraph(); r = par.add_run("Aturan Dua Orang Berbasis AI untuk Melindungi Lansia dari Penipuan Digital")
    set_run_font(r, 16, bold=True)
    fmt(par, spacing=1.15, after=18, align=WD_ALIGN_PARAGRAPH.CENTER)
    for line in ["Tema: Human-Centered Engineering in the AI Era", f"Subtema: {cfg['subtema']}"]:
        par = d.add_paragraph(); r = par.add_run(line); set_run_font(r, 12, italic=True)
        fmt(par, spacing=1.15, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    fmt(d.add_paragraph(), spacing=1.0, after=40)
    par = d.add_paragraph(); r = par.add_run(f"Disusun oleh Tim {cfg['nama_tim']}"); set_run_font(r, 13, bold=True)
    fmt(par, spacing=1.15, after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    for m in cfg["anggota"]:
        par = d.add_paragraph()
        r = par.add_run(f"{m['nama']} ({m['peran']}) - {m['prodi']}, {m['kampus']}"); set_run_font(r, 12)
        fmt(par, spacing=1.15, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    fmt(d.add_paragraph(), spacing=1.0, after=60)
    for line in ["Hackvidia - Arkavidia 11.0", "Himpunan Mahasiswa Informatika Institut Teknologi Bandung", "2026"]:
        par = d.add_paragraph(); r = par.add_run(line); set_run_font(r, 12, bold=True)
        fmt(par, spacing=1.15, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)

    # new section for main text with page numbers from 1
    sec = d.add_section(WD_SECTION.NEW_PAGE)
    restart_numbering(sec, 1)
    add_page_number_footer(sec)

    v = cfg["validasi"]

    # ---------------- 1
    D.h1("1. Ringkasan Eksekutif")
    D.p("Penipuan digital paling sering berhasil pada satu momen yang sama: ketika seseorang diminta uang secara mendesak "
        "dan sedang sendirian. Indonesia Anti-Scam Centre (IASC) OJK menerima **702.342 laporan penipuan** sejak November 2024 "
        "hingga September 2026 [1]. Modus *fake call*, yaitu penipu yang menyamar sebagai anggota keluarga, polisi, atau petugas "
        "lembaga, menimbulkan kerugian terlapor terbesar, **Rp2,09 triliun** [3]. Lansia berada di posisi paling rentan: indeks "
        "literasi keuangan kelompok umur 51-79 tahun hanya 54,55% [9], sementara lansia yang mengakses internet naik dari 14,10% "
        "(2021) menjadi 34,13% (2025) [8]. Riset McAfee menunjukkan tiga detik rekaman sudah cukup untuk membuat tiruan suara "
        "dengan kecocokan sekitar 85% [6].")
    D.p("**JagaOrtu** membawa prinsip yang lama dipakai bank untuk transaksi berisiko, yaitu *aturan dua orang* (maker-checker), "
        "ke kehidupan sehari-hari lansia. Ketika pesan WhatsApp/SMS, telepon, atau cerita lansia menunjukkan tanda penipuan, JagaOrtu "
        "(1) menghentikan sejenak dan menjelaskan tanda-tandanya dengan suara dalam bahasa lansia (Indonesia, Sunda, atau Jawa), "
        "(2) menghadirkan anggota keluarga yang ditunjuk (Pendamping) dalam satu ketukan, dan (3) memandu verifikasi dengan kata "
        "sandi keluarga, cara yang tetap ampuh ketika suara dapat dipalsukan. Fitur *Latihan Waspada* membangun kewaspadaan lewat "
        "simulasi yang aman.")
    D.p("AI dipakai secukupnya: aturan pasti dan model di perangkat menangani sebagian besar kasus, model bahasa besar hanya "
        "dipanggil untuk kasus ambigu dengan data yang telah disamarkan, dan keputusan akhir selalu berada di tangan manusia. "
        "JagaOrtu dimonetisasi melalui model B2B2C bagi bank dan BPR yang wajib menjalankan strategi anti-fraud "
        "(POJK 12/2024) [11] serta dompet digital, ditambah paket keluarga. Pilot dimulai di Bandung karena Jawa Barat mencatat laporan penipuan terbanyak "
        "secara nasional, yaitu 131.445 laporan [3].")

    # ---------------- 2
    D.h1("2. Latar Belakang dan Analisis Masalah")
    D.h2("2.1 Skala dan urgensi")
    D.p("IASC mencatat kerugian Rp9,1 triliun dari 432.637 aduan hingga 14 Januari 2026 [2], atau rata-rata sekitar Rp21 juta "
        "per aduan (perhitungan tim). Tiga modus teratas nasional adalah belanja daring (113.520 laporan; Rp1,78 triliun), "
        "*fake call* (63.517 laporan; Rp2,09 triliun), dan investasi (32.251 laporan; Rp2,08 triliun) [3]. Ancaman ini makin "
        "canggih: laporan penipuan yang terindikasi memakai AI berjumlah 13 pada akhir 2024 dan bertambah menjadi total 1.379 hingga Juni 2026 [4], dan "
        "Kementerian Komdigi menyebut banyak lansia menjadi korban penipuan yang meniru suara pejabat dengan AI [5]. Pola ini "
        "global: di Amerika Serikat, kelompok usia 60+ menanggung kerugian terbesar, US$4,885 miliar dari 147.127 aduan pada "
        "2024 [13].")
    D.p("Kelompok yang terpapar terus membesar. SUPAS 2025 mencatat porsi lansia 11,97% dari 284,67 juta penduduk [7], atau "
        "sekitar 34 juta orang (perhitungan tim). Sebanyak 52,23% lansia telah menggunakan telepon seluler dan 34,13% mengakses "
        "internet [8]. Artinya, semakin banyak lansia masuk ke ruang digital lebih cepat daripada kesiapan literasinya.")
    D.h2("2.2 Akar masalah: mengapa lansia tertipu")
    D.p("Analisis terhadap modus yang dilaporkan menunjukkan tiga mekanisme yang berulang. **Pertama, urgensi dan otoritas:** "
        "penipu menciptakan kepanikan (\"anak kecelakaan\", \"anak ditahan polisi\", \"rekening akan diblokir\") dan menyamar "
        "sebagai pihak yang dipercaya. **Kedua, isolasi:** korban diminta merahasiakan dan segera bertindak, sehingga keputusan "
        "diambil sendirian. **Ketiga, runtuhnya sinyal kepercayaan lama:** suara, foto, bahkan nomor yang dikenal kini dapat "
        "dipalsukan [6]. Implikasinya, yang dibutuhkan lansia bukan sekadar informasi tambahan, melainkan **jeda dan kehadiran "
        "orang kedua yang dipercaya tepat pada momen kritis**. Edukasi penting, tetapi tidak hadir saat lansia sedang ditekan "
        "penipu.")
    D.h2("2.3 Pengguna yang kami rancang")
    D.p("**Bu Rukmini (67)**, pensiunan guru di Bandung, tinggal sendiri, memakai WhatsApp dan *mobile banking* untuk menerima "
        "pensiun. Ia mudah panik bila mendengar kabar buruk tentang keluarga dan sungkan bertanya karena takut merepotkan. "
        "**Rina (34)**, anaknya yang bekerja di Jakarta, ingin melindungi ibunya tanpa memata-matai isi percakapannya. Solusi "
        "harus menghormati martabat dan otonomi Bu Rukmini sekaligus memberi Rina cara untuk hadir saat dibutuhkan.")
    D.h2("2.4 Celah solusi yang ada")
    D.table(["Solusi saat ini", "Yang sudah dilakukan", "Celah bagi lansia"], [
        ["Edukasi OJK/Komdigi", "Meningkatkan kesadaran masyarakat", "Tidak hadir pada momen serangan"],
        ["Aplikasi caller ID (Truecaller, Getcontact)", "Menandai nomor spam", "Nomor penipu cepat berganti; tidak membaca isi pesan WhatsApp, tautan, atau file .apk; tidak melibatkan keluarga"],
        ["Google Fake Call Detection (2026)", "Memperingatkan bila penelepon memalsukan identitas kontak tersimpan", "Bergantung pada aplikasi Phone by Google (Android 12+, mulai di Pixel) [10]; tidak mencakup WhatsApp; tidak ada orang kedua"],
        ["Cek rekening (cekrekening.id)", "Mengecek rekening terlapor", "Harus dibuka manual; jarang diketahui lansia"],
        ["Peringatan aplikasi bank", "Muncul saat transfer", "Generik, mudah diabaikan, tanpa penjelasan bersuara"],
    ], [4.2, 4.4, 7.3], "Perbandingan solusi yang ada dan celahnya")
    D.p("Belum ada lapisan yang sekaligus (1) menangkap momen berisiko secara otomatis lintas kanal, termasuk WhatsApp, "
        "(2) menghadirkan orang kedua yang dipercaya tepat pada momen itu, (3) berbicara dalam bahasa lansia, dan (4) melatih "
        "kewaspadaan secara berkala.")
    D.p(validation_paragraph(v))

    # ---------------- 3
    D.h1("3. Deskripsi Solusi")
    D.p("JagaOrtu terdiri atas aplikasi Android untuk lansia dan aplikasi Pendamping untuk keluarga. Lima fitur intinya "
        "dirancang mengikuti perjalanan lansia saat menghadapi penipuan (Tabel 2).")
    D.table(["Fitur", "Cara kerja", "Nilai bagi pengguna"], [
        ["**Penjaga Otomatis**", "Membaca notifikasi WhatsApp/SMS dari nomor tak dikenal atau yang berisi uang, tautan, dan file, lalu menilai risikonya di perangkat", "Lansia tidak perlu curiga lebih dulu; peringatan datang sebelum uang berpindah"],
        ["**Tanya Dulu**", "Satu tombol besar untuk bercerita lewat suara atau membagikan pesan dan tangkapan layar", "Bertanya terasa wajar, tidak memalukan"],
        ["**Aturan Dua Orang**", "Satu ketukan mengirim ringkasan risiko ke Pendamping; tanpa respons 10 menit, pendamping cadangan dihubungi", "Keputusan tidak lagi diambil sendirian"],
        ["**Verifikasi Keluarga**", "Saat telepon mengaku keluarga atau polisi, JagaOrtu memandu: tanyakan kata sandi keluarga, tutup, lalu telepon balik nomor tersimpan", "Tetap ampuh walau suara dikloning AI"],
        ["**Latihan Waspada**", "Simulasi pesan penipuan yang aman atas izin keluarga, diikuti pelajaran suara 30 detik", "Kewaspadaan terbentuk sebelum penipuan sungguhan datang"],
    ], [3.3, 7.3, 5.3], "Fitur inti JagaOrtu")
    D.figure(BUILD / "journey.png", 6.27, "Perjalanan pengguna tanpa dan dengan JagaOrtu")
    D.p("Rancangan ini berpegang pada lima prinsip *human-centered*: **martabat** (peringatan menjelaskan tanda bahaya tanpa "
        "menghakimi), **otonomi** (lansia tetap memutuskan; JagaOrtu tidak pernah memblokir atau mentransfer sendiri), **suara "
        "dan bahasa ibu**, **kesederhanaan** (maksimal dua aksi utama per layar), dan **privasi** (Pendamping hanya melihat "
        "ringkasan risiko, bukan isi chat). Pembeda utamanya adalah cara kerja yang meniru kebiasaan manusia yang paling efektif "
        "melawan penipuan, yaitu berhenti sejenak dan menelepon orang yang dipercaya. Fitur Latihan Waspada menerapkan prinsip "
        "*inoculation*: paparan ringan yang aman terbukti meningkatkan ketahanan terhadap manipulasi [17].")

    # ---------------- 4
    D.h1("4. Implementasi AI")
    D.h2("4.1 Prinsip \"AI secukupnya\"")
    D.p("Tema Hackvidia menuntut kejelasan kapan, bagaimana, dan sejauh mana AI diperlukan. JagaOrtu menjawabnya dengan empat "
        "lapis keputusan (Tabel 3). Setiap lapis hanya dipakai bila lapis sebelumnya tidak cukup.")
    D.table(["Lapis", "Teknologi", "Kapan dipakai", "Alasan"], [
        ["L0 Aturan pasti", "Daftar nomor/rekening/domain terlapor, deteksi file .apk, permintaan OTP/PIN, kata kunci transfer mendesak", "Semua pesan dari nomor tak dikenal atau yang berisi uang, tautan, dan file", "Cepat, gratis, dapat dijelaskan"],
        ["L1 Model di perangkat", "Klasifikasi teks ringan (TensorFlow Lite) dan OCR on-device", "Pesan lolos L0 tetapi berpola mencurigakan", "Privasi terjaga, bekerja offline, kurang dari 1 detik"],
        ["L2 Model bahasa besar", "LLM dengan keluaran terstruktur: modus, tanda bahaya, penjelasan, saran", "Hanya skor ambigu; teks disamarkan lebih dulu", "Memahami konteks baru, bahasa campur dan bahasa daerah"],
        ["L3 Manusia", "Lansia dan Pendamping", "Setiap keputusan uang", "AI membantu menimbang, manusia memutuskan"],
    ], [2.6, 5.0, 4.3, 4.0], "Empat lapis keputusan JagaOrtu")
    D.h2("4.2 Domain AI dan alasan pemilihan")
    D.bullets([
        "**Natural Language Processing:** pesan penipuan berbentuk teks bebas yang bercampur bahasa Indonesia, bahasa daerah, dan "
        "singkatan. Klasifikasi modus serta penjelasan sederhana memerlukan pemahaman konteks, bukan sekadar kata kunci.",
        "**Speech (speech-to-text dan text-to-speech):** lansia lebih nyaman bercerita daripada mengetik dan lebih mudah memahami "
        "penjelasan yang didengar. Layanan TTS neural sudah mendukung suara Indonesia, Sunda, dan Jawa [16].",
        "**Computer Vision (OCR):** bukti transfer palsu, kode QR, dan tangkapan layar tautan perlu dibaca agar dapat dicek oleh "
        "aturan L0 dan model L1.",
    ])
    D.p("JagaOrtu sengaja **tidak mengklaim mendeteksi deepfake suara**. Akurasi deteksi kloning suara belum andal dan cepat "
        "usang, sedangkan verifikasi prosedural (kata sandi keluarga dan telepon balik) tetap efektif apa pun kualitas suaranya. "
        "Inilah keputusan desain yang menempatkan AI hanya di bagian yang benar-benar memberi nilai.")
    D.h2("4.3 Arsitektur sistem")
    D.figure(BUILD / "arch.png", 6.27, "Arsitektur JagaOrtu dan alur keputusan empat lapis")
    D.p("Di perangkat lansia, *NotificationListenerService* membaca notifikasi atas izin eksplisit, *CallScreeningService* "
        "mengenali panggilan masuk tanpa merekam isi percakapan [19], dan tautan dicek melalui Google Safe Browsing API. Layanan "
        "cloud hanya menerima teks yang sudah disamarkan, lalu mengembalikan modus, penjelasan, dan audio TTS. Aplikasi Pendamping "
        "menerima ringkasan risiko dan aksi satu ketukan.")
    D.h2("4.4 Keamanan dan privasi")
    D.p("Prinsip minimisasi data mengikuti UU No. 27 Tahun 2022 tentang Pelindungan Data Pribadi [12]: pemrosesan di perangkat "
        "lebih dulu; nomor, nama, dan rekening disamarkan sebelum dikirim ke cloud; isi pesan tidak disimpan; Pendamping hanya "
        "menerima ringkasan; lansia dapat menjeda perlindungan atau menghapus riwayat kapan saja. Kontak keluarga tidak dibaca "
        "secara bawaan.")
    D.h2("4.5 Batasan platform dan mitigasinya")
    D.p("Kami merancang JagaOrtu dengan jujur terhadap batas Android. **(1) OTP:** sejak Android 15, aplikasi "
        "*NotificationListenerService* yang tidak berstatus tepercaya tidak dapat membaca isi notifikasi yang terdeteksi "
        "memuat OTP [18]. JagaOrtu memang tidak membutuhkannya: yang dinilai adalah pesan yang *meminta* kode, tautan, atau "
        "transfer, dan pesan itu tetap terbaca. **(2) WhatsApp:** pesan terenkripsi *end-to-end* [20], sehingga JagaOrtu hanya "
        "melihat teks yang ditampilkan notifikasi (nama pengirim dan potongan pesan), bukan isi chat, foto, atau *voice note*. "
        "Bila pratinjau dimatikan atau pesan terpotong, lansia diarahkan ke Tanya Dulu untuk membagikan pesan secara sukarela. "
        "**(3) Telepon:** *CallScreeningService* hanya aktif bila pengguna memberi peran ROLE_CALL_SCREENING, hanya menerima "
        "nomor dan metadata panggilan seluler (tanpa audio), harus merespons dalam 5 detik, dan tanpa izin kontak hanya menerima "
        "panggilan dari nomor yang tidak tersimpan [19]. Panggilan WhatsApp tidak melewati layanan ini, sehingga JagaOrtu "
        "mengandalkan notifikasi panggilan masuk dan pengingat kata sandi keluarga. **(4) Pemasangan:** aplikasi di luar Play "
        "Store perlu \"Izinkan setelan terbatas\" sebelum akses notifikasi dapat diaktifkan [21]; langkah ini masuk ke "
        "panduan pemasangan oleh anak.")

    # ---------------- 5
    D.h1("5. Desain UX dan Wireframe")
    D.p("Desain mengikuti kebutuhan aksesibilitas lansia: teks minimal 20 sp, kontras tinggi, tombol besar, satu aksi utama yang "
        "jelas, peringatan yang otomatis dibacakan, dan nada yang menenangkan. Tiga alur utama digambarkan pada Lampiran A:")
    D.bullets([
        "**Pemasangan oleh anak** (Gambar A1): tiga langkah berupa izin terbatas, pemilihan Pendamping utama dan cadangan, serta "
        "kata sandi keluarga dan bahasa suara.",
        "**Momen berisiko** (Gambar A3-A6): pesan WhatsApp \"anak minta transfer\" memicu peringatan bersuara; lansia menekan "
        "\"Tanya Rina dulu\"; Rina menerima ringkasan dan menelepon. Untuk telepon, JagaOrtu mengingatkan kata sandi keluarga.",
        "**Penguatan jangka panjang** (Gambar A7-A8): Latihan Waspada dan pengaturan privasi yang dikendalikan lansia.",
    ], numbered=True)

    # ---------------- 6
    D.h1("6. Rencana Data")
    D.p("**Sumber data.** (1) Dataset SMS berbahasa Indonesia berisi 1.143 pesan berlabel normal, penipuan, dan promo [15]; "
        "(2) contoh modus dari peringatan resmi OJK, Komdigi, BSSN, bank, dan pemberitaan, ditranskripsi dan dilabeli ulang; "
        "(3) variasi sintetis yang dihasilkan LLM dari pola nyata, termasuk bahasa Sunda dan Jawa, lalu ditinjau manusia; serta "
        "(4) laporan pengguna yang bersifat opt-in dan anonim untuk pembaruan berkelanjutan.")
    D.p("**Taksonomi.** Delapan kelas yang selaras dengan kategori laporan IASC: keluarga darurat, menyamar petugas (polisi, "
        "bank, BPJS, CS), tautan atau file berbahaya, hadiah atau undian, investasi atau *deepfake* tokoh, salah transfer atau "
        "refund, belanja atau COD fiktif, dan pinjaman ilegal. Ditambah satu kelas aman.")
    D.p("**Pelabelan dan evaluasi.** Setiap sampel dilabeli dua anotator dengan adjudikasi (target Cohen's kappa 0,8 atau "
        "lebih). Model dinilai pada data nyata yang ditahan dan tidak dipakai saat pelatihan, dengan target recall 90% atau lebih "
        "untuk kelas berisiko tinggi dan alarm palsu 10% atau kurang pada pesan normal. Untuk aspek manusia, diukur pula waktu "
        "hingga keputusan aman dan tingkat pemahaman lansia terhadap penjelasan (*teach-back*). Data pribadi tidak disimpan; "
        "laporan anonim dihapus setelah 30 hari kecuali disetujui untuk korpus.")

    # ---------------- 7
    D.h1("7. Rancangan Implementasi, Kelayakan, dan Skalabilitas")
    D.p("**Tech stack.** Android (Kotlin) dengan TensorFlow Lite dan ML Kit OCR; backend FastAPI di layanan serverless; "
        "PostgreSQL untuk basis laporan dan taksonomi; API LLM dengan keluaran JSON terstruktur; layanan STT/TTS neural; "
        "aplikasi Pendamping berbasis web progresif dengan notifikasi push.")
    D.table(["Fase", "Waktu", "Target"], [
        ["0. MVP", "Hacking Day (24 jam)", "Penjaga Otomatis untuk WhatsApp/SMS, peringatan bersuara, Tanya Dulu, notifikasi Pendamping, demo langsung"],
        ["1. Pilot", "Februari-April 2027", "30 keluarga di Bandung; mengukur alarm palsu, waktu check-in, dan pemahaman lansia"],
        ["2. Komunitas", "Mei-Desember 2027", "1.000 keluarga melalui posyandu lansia dan pengajian; Latihan Waspada rutin"],
        ["3. Mitra", "2028", "API \"Jeda Transfer\" dengan bank, BPR, dan e-wallet sebagai bagian strategi anti-fraud"],
    ], [2.6, 3.8, 9.5], "Peta jalan pengembangan")
    D.p("**Kelayakan.** Seluruh komponen memakai API Android resmi dan model atau layanan yang sudah tersedia, sehingga tidak "
        "perlu melatih model besar dari nol. Ruang lingkup MVP dipilih agar dapat didemonstrasikan secara langsung pada akhir "
        "Hacking Day. **Skalabilitas.** Karena sebagian besar keputusan diambil di perangkat, biaya cloud per pengguna tetap "
        "rendah ketika pengguna bertambah. Taksonomi dan aturan diperbarui terpusat setiap minggu tanpa memperbarui aplikasi.")
    D.p("**Risiko utama dan mitigasinya:** alarm palsu (ambang yang dapat disetel dan umpan balik \"ini aman\"); kesulitan "
        "pemasangan (dipasang oleh anak dalam tiga langkah); keterbatasan iOS (dimulai dari Android, yang dipakai mayoritas "
        "pengguna ponsel di Indonesia); serta lansia tanpa keluarga dekat (pendamping cadangan dari kader komunitas).")

    # ---------------- 8
    D.h1("8. Dampak, Model Bisnis, dan Keberlanjutan")
    D.h2("8.1 Dampak")
    D.p("Pasar sasaran adalah sekitar 34 juta lansia, dengan lebih dari sepertiganya sudah daring [7][8]. Dengan rata-rata "
        "kerugian sekitar Rp21 juta per aduan [2], satu penipuan yang dicegah setara dengan tabungan berbulan-bulan seorang "
        "pensiunan. Skenario konservatif (asumsi untuk divalidasi saat pilot): jika JagaOrtu mencegah satu kejadian per 50 "
        "keluarga per tahun, maka 1.000 keluarga menghindari kerugian sekitar Rp420 juta per tahun, dan 100.000 pengguna melalui "
        "mitra bank sekitar Rp42 miliar. Dampak nonfinansialnya: lansia lebih percaya diri memakai layanan digital dan keluarga "
        "lebih terhubung.")
    D.h2("8.2 Model bisnis")
    D.p("Model bisnis konsisten dengan Lean Canvas. **Pelanggan B2B2C:** bank dan BPR yang wajib menerapkan "
        "strategi anti-fraud empat pilar (pencegahan, deteksi, investigasi, pemantauan) menurut POJK 12/2024 [11], serta dompet digital, membayar lisensi per pengguna terlindungi "
        "(asumsi awal Rp4.000 per pengguna per bulan). **Paket Keluarga:** perlindungan dasar gratis; premium Rp25.000 per bulan "
        "(asumsi) untuk dua lansia, tiga pendamping, dan Latihan Waspada. **Pendanaan tambahan:** hibah atau CSR program literasi "
        "digital lansia. **Struktur biaya:** dengan model kecil (sekitar 1.000 token per analisis), biaya LLM di bawah Rp10 per "
        "analisis, atau di bawah Rp1.000 per pengguna per bulan untuk maksimal 90 analisis cloud, ditambah TTS, server, dan "
        "pengembangan (estimasi). **Kanal:** anak memasang di ponsel orang tua, komunitas lansia, dan kemitraan "
        "lembaga keuangan.")
    D.h2("8.3 Indikator keberhasilan dan keberlanjutan")
    D.p("Indikator utama: lansia aktif mingguan, momen berisiko yang tertangkap, check-in Pendamping yang selesai dalam 10 "
        "menit, tingkat alarm palsu, tingkat lulus Latihan Waspada, dan retensi keluarga tiga bulan. Keberlanjutan dijaga oleh "
        "pendapatan B2B2C yang didorong regulasi, efek jaringan keluarga, dan korpus modus lokal yang terus tumbuh. JagaOrtu "
        "mewujudkan *Human-Centered Engineering in the AI Era*: AI dipakai untuk memahami dan menjelaskan, sedangkan kendali tetap "
        "berada pada lansia dan orang-orang yang mereka percaya.")

    # ---------------- 9
    D.h1("9. Penutup")
    D.p("Penipu menang ketika korbannya panik dan sendirian. JagaOrtu memutus kedua kondisi itu dengan jeda yang tepat waktu, "
        "penjelasan dalam bahasa yang dipahami lansia, dan kehadiran keluarga dalam satu ketukan. Tujuan kami sederhana: tidak "
        "ada lansia yang harus menghadapi penipu sendirian.")

    # ---------------- Daftar Pustaka
    D.h1("Daftar Pustaka")
    for i, ref in enumerate(REFS, 1):
        par = d.add_paragraph()
        r = par.add_run(f"[{i}] {ref}"); set_run_font(r, 12)
        pf = par.paragraph_format
        pf.left_indent = Cm(0.9); pf.first_line_indent = Cm(-0.9)
        fmt(par, spacing=1.0, after=4, align=WD_ALIGN_PARAGRAPH.LEFT)

    # ---------------- Lampiran (tidak dihitung dalam batas 10 halaman)
    sec2 = d.add_section(WD_SECTION.NEW_PAGE)
    for el in sec2._sectPr.findall(qn("w:pgNumType")):
        sec2._sectPr.remove(el)
    for i, name in enumerate(["Lampiran_A1", "Lampiran_A2", "Lampiran_A3", "Lampiran_A4"]):
        if i == 0:
            D.h1("Lampiran A. Wireframe Low-Fidelity")
        img = trim(ROOT / "03_Wireframe" / f"{name}.png")
        w, h = Image.open(img).size
        width = min(6.27, 8.9 * w / h)
        par = d.add_paragraph(); par.add_run().add_picture(str(img), width=Inches(width))
        fmt(par, spacing=1.0, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
        if i < 3:
            D.page_break()
    return d


def check_pages(pdf_path, limit=10):
    """Hitung halaman teks utama (setelah cover, sebelum 'LAMPIRAN A') dan beri peringatan bila > limit."""
    try:
        from pypdf import PdfReader
    except ImportError:
        print("  (lewati cek halaman: pypdf belum terpasang)")
        return None
    pages = PdfReader(str(pdf_path)).pages
    first_lamp = next((i for i, pg in enumerate(pages) if "LAMPIRAN A" in (pg.extract_text() or "")), len(pages))
    n = first_lamp - 1  # halaman 0 = cover
    status = "OK" if n <= limit else "MELEBIHI BATAS"
    print(f"  halaman teks: {n} (batas {limit}) -> {status}; total PDF {len(pages)} halaman")
    return n


def main():
    cfg = load_config()
    d = build(cfg)
    out_dir = ROOT / "02_Proposal"
    docx_path = out_dir / f"{cfg['nama_tim']}_Proposal.docx"
    pdf_path = out_dir / f"{cfg['nama_tim']}_Proposal.pdf"
    d.save(docx_path)
    if docx_to_pdf(docx_path, pdf_path):
        print("proposal ok ->", pdf_path)
        check_pages(pdf_path)


if __name__ == "__main__":
    main()
